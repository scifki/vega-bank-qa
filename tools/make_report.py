from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

root = Path(__file__).resolve().parent.parent
shots = root / "screenshots"

doc = Document()

# ---------- стили по ГОСТ 7.32 ----------
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(14)
normal.font.color.rgb = RGBColor(0, 0, 0)
pf = normal.paragraph_format
pf.line_spacing = 1.5
pf.space_after = Pt(0)
pf.space_before = Pt(0)
pf.first_line_indent = Cm(1.25)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(3)
section.right_margin = Cm(1.5)

TABLES = 0
FIGURES = 0
LISTINGS = 0


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


section.different_first_page_header_footer = True
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.paragraph_format.first_line_indent = Cm(0)
add_page_number(footer_p)


def h(text, level=1):
    par = doc.add_heading(level=level)
    par.paragraph_format.first_line_indent = Cm(1.25 if level == 2 else 0)
    par.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    par.paragraph_format.space_after = Pt(6)
    par.paragraph_format.line_spacing = 1.5
    par.paragraph_format.keep_with_next = True
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    run = par.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return par


def p(text="", bold=False, italic=False, align=None, indent=Cm(1.25), size=14):
    par = doc.add_paragraph()
    par.paragraph_format.first_line_indent = indent
    if align is not None:
        par.alignment = align
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return par


def bullets(items):
    for item in items:
        par = doc.add_paragraph(item)
        par.paragraph_format.first_line_indent = Cm(1.25)
        par.paragraph_format.left_indent = Cm(0.5)
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def shade(paragraph, color="F2F2F2"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def code(text, caption=None):
    global LISTINGS
    if caption:
        LISTINGS += 1
        cp = doc.add_paragraph()
        cp.paragraph_format.first_line_indent = Cm(0)
        cp.paragraph_format.space_before = Pt(6)
        cp.paragraph_format.keep_with_next = True
        run = cp.add_run(f"Листинг {LISTINGS} – {caption}")
        run.font.size = Pt(12)
    par = doc.add_paragraph()
    par.paragraph_format.first_line_indent = Cm(0)
    par.paragraph_format.space_after = Pt(10)
    par.paragraph_format.line_spacing = 1.0
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for i, line in enumerate(text.strip("\n").split("\n")):
        run = par.add_run(("" if i == 0 else "\n") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    shade(par)


def plain_table(rows, widths=None, header=None):
    cols = len(rows[0])
    t = doc.add_table(rows=0, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if header:
        cells = t.add_row().cells
        for i, text in enumerate(header):
            par = cells[i].paragraphs[0]
            par.paragraph_format.first_line_indent = Cm(0)
            run = par.add_run(text)
            run.bold = True
            run.font.size = Pt(12)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            par = cells[i].paragraphs[0]
            par.paragraph_format.first_line_indent = Cm(0)
            run = par.add_run(str(value))
            run.font.size = Pt(12)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w
    return t


def table(caption, headers, rows, widths=None):
    global TABLES
    TABLES += 1
    cap = doc.add_paragraph()
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_before = Pt(10)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.keep_with_next = True
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = cap.add_run(f"Таблица {TABLES} – {caption}")
    run.font.size = Pt(12)

    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0].cells
    for i, text in enumerate(headers):
        par = hdr[i].paragraphs[0]
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.line_spacing = 1.0
        run = par.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        shade(par, "E7E9EE")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            par = cells[i].paragraphs[0]
            par.paragraph_format.first_line_indent = Cm(0)
            par.paragraph_format.line_spacing = 1.0
            run = par.add_run(str(value))
            run.font.size = Pt(11)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def img(name, caption, width=15.5):
    global FIGURES
    FIGURES += 1
    par = doc.add_paragraph()
    par.paragraph_format.first_line_indent = Cm(0)
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.keep_with_next = True
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(shots / name), width=Cm(width))

    cap = doc.add_paragraph()
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_after = Pt(10)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(f"Рисунок {FIGURES} – {caption}")
    run.font.size = Pt(12)


# ---------- титульный лист ----------
for _ in range(4):
    doc.add_paragraph()

tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.paragraph_format.first_line_indent = Cm(0)
r = tp.add_run("ПРАКТИЧЕСКАЯ РАБОТА")
r.bold = True
r.font.size = Pt(20)

tp2 = doc.add_paragraph()
tp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp2.paragraph_format.first_line_indent = Cm(0)
r = tp2.add_run("Python: организация автоматизации\nтестирования с нуля")
r.bold = True
r.font.size = Pt(16)

tp3 = doc.add_paragraph()
tp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp3.paragraph_format.first_line_indent = Cm(0)
r = tp3.add_run("Учебный кейс «Вега-Банк»")
r.italic = True
r.font.size = Pt(14)

for _ in range(5):
    doc.add_paragraph()

plain_table(
    [
        ("Выполнил(а)", "Лисевич Арсений"),
        ("Группа", "Исп/п23-09.2"),
        ("Дисциплина", "Поддержка и тестирование программных модулей"),
        ("Проверил(а)", "________________________________"),
        ("Дата", "28.10.2026"),
    ],
    widths=[Cm(4.2), Cm(10.3)],
)

doc.add_paragraph()
doc.add_paragraph()
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
end.paragraph_format.first_line_indent = Cm(0)
r = end.add_run("2026")
r.font.size = Pt(14)

doc.add_page_break()

# ---------- содержание ----------
h("СОДЕРЖАНИЕ", 1)
contents = [
    "Задание 1. Изучаем проект",
    "Задание 2. Требования и стек на Python",
    "Задание 3. Инфраструктура и каркас репозитория",
    "Задание 4. Якоря на UI и сценарий Gherkin",
    "Задание 5. Шаги на Python, запуск и отчёт",
    "Задание 6. Демо и план внедрения",
    "Контрольные вопросы",
    "Структура проекта",
    "Как запустить",
]
for item in contents:
    par = doc.add_paragraph(item)
    par.paragraph_format.first_line_indent = Cm(0)
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT

doc.add_page_break()

# ---------- Задание 1 ----------
h("Задание 1. Изучаем проект", 1)

h("1.1. Для кого создан продукт и на какие проверки сделать упор", 2)
p("Продукт кейса — онлайн-банк для физических лиц: переводы между своими счетами, платежи, "
  "карты, кредиты, личный кабинет и уведомления. Пользователи — обычные клиенты банка, для которых "
  "главное — быстро и без ошибок провести денежную операцию и видеть корректный остаток.")
p("Поэтому упор в автоматизации делается на критичные денежные сценарии (платежи и переводы, frontend + backend) "
  "и на регресс UI личного кабинета: проверяются позитивные и негативные суммы (0, отрицательная, больше остатка, "
  "граничные значения), тексты сообщений об ошибках и целостность данных в базе после операции.")

h("1.2. Участники разработки", 2)
table(
    "Участники разработки",
    ["Участник", "Роль для автоматизации", "Что нужно от него получить"],
    [
        [
            "Владельцы продуктов",
            "Заказчики: определяют ценность и приоритет функционала, решают, что автоматизировать первым",
            "Приоритизированный бэклог, критерии приёмки, согласованный объём регресса и даты демо",
        ],
        [
            "QA каждой команды",
            "Клиенты инструмента автоматизации: знают ручные проверки и данные, принимают автотесты",
            "Тест-кейсы и чек-листы, тестовые данные, обратную связь по падениям, приоритеты регресса",
        ],
        [
            "Лидер ручного тестирования",
            "Держит единый стандарт качества и регресс-план, обучает QA",
            "Реестр регресса с приоритетами, правила оформления кейсов, план обучения команд",
        ],
        [
            "Лидер разработки Frontend",
            "Отвечает за разметку, от которой зависят локаторы автотестов",
            "Обязательные data-test-id в новых фичах, стабильные локаторы, доступ к стендам, ревью правок UI",
        ],
        [
            "Специалист по закупке",
            "Обеспечивает инфраструктуру: серверы, лицензии, доступы",
            "Заявку на VM под CI-раннеры и сроки поставки, лицензии, согласование бюджета",
        ],
    ],
    widths=[Cm(3.6), Cm(5.6), Cm(5.9)],
)

h("1.3. Три вопроса команде «Платежи»", 2)
bullets(
    [
        "1. Какие сценарии вы проверяете вручную чаще всего и какие из них критичны для клиента "
        "(перевод, лимиты, ошибки ввода)? — чтобы выбрать первые автотесты по реальному риску, а не «на глаз».",
        "2. Какие части платежей меняются в ближайшие спринты, а какие API и формы уже стабильны? — "
        "нестабильное автоматизировать рано, стабильное даёт быстрый эффект.",
        "3. Есть ли постоянные тестовые счета, стенд и доступ к базе, чтобы проверять остатки после перевода? — "
        "без данных и доступа к БД автотест проверит только текст на экране.",
    ]
)

# ---------- Задание 2 ----------
h("Задание 2. Требования и стек на Python", 1)

h("2.1. Выбор инструментов", 2)
table(
    "Выбор инструментов автоматизации",
    ["Задача", "Варианты", "Выбор и обоснование"],
    [
        [
            "Тестирование Frontend",
            "Selenium, Playwright",
            "Playwright: авто-ожидания без sleep, один API для Chromium, Firefox и WebKit, "
            "трассировка и кодогенерация — ручному QA проще освоить",
        ],
        [
            "Тестирование Backend (REST)",
            "requests, httpx",
            "httpx: тот же синтаксис, что у requests, но есть таймауты по умолчанию, "
            "поддержка async и HTTP/2 — запас на рост нагрузки",
        ],
        [
            "Обращения в базу данных",
            "sqlite3, psycopg2, SQLAlchemy",
            "SQLAlchemy Core с драйвером psycopg2: один код для локальной sqlite и тестового PostgreSQL, "
            "параметризованные запросы без склейки SQL",
        ],
        [
            "Тесты пишут ручные QA без опыта",
            "behave, pytest-bdd",
            "behave: сценарии на Gherkin по-русски, шаги переиспользуются, "
            "не нужно знать фикстуры pytest — ниже порог входа",
        ],
        [
            "Отчёты о прогоне",
            "Allure, pytest-html",
            "HTML-отчёт, который собирается из JSON-результатов behave скриптом make_html_report.py: "
            "один самодостаточный файл без внешних серверов и Java, открывается в браузере и прикладывается к CI",
        ],
    ],
    widths=[Cm(3.4), Cm(3.2), Cm(8.5)],
)

h("2.2. Преимущество и риск Python-стека", 2)
p("Преимущество: низкий порог входа — ручные тестировщики читают Gherkin-сценарии и пишут шаги на Python "
  "быстрее, чем на Java, а библиотеки Playwright и httpx ставятся одной командой, поэтому первые автотесты "
  "появляются в первом же спринте.")
p("Риск: Python-стек в крупной промышленной эксплуатации требует самодельной обвязки (параллельный запуск, "
  "распределённые прогоны, стабильность тестов), и готовых решений «из коробки» меньше, чем в Java "
  "(Selenium Grid, TestNG, Maven). Если каркас не поддерживать, тесты начнут падать и потеряют доверие.")

h("2.3. Сколько нужно автоматизаторов", 2)
p("Норма: один автоматизатор обслуживает не больше пяти команд. Всего команд — 6.")
p("Расчёт: 6 / 5 = 1,2 → округляем вверх = 2 автоматизатора. Второй нужен не «на всякий случай»: "
  "пять команд закрывает первый, шестая плюс обучение QA и поддержка каркаса — второй; "
  "он же подменяет первого в отпуске и болезни.")

# ---------- Задание 3 ----------
h("Задание 3. Инфраструктура и каркас репозитория", 1)

h("3.1. Машины и сервисы для удалённого параллельного запуска", 2)
table(
    "Инфраструктура запуска автотестов",
    ["Задача", "Варианты", "Выбор и обоснование"],
    [
        [
            "Сервер CI",
            "GitHub Actions, GitLab CI, Jenkins",
            "GitHub Actions: репозиторий уже в GitHub, пайплайн описывается файлом tests.yml, "
            "бесплатные минуты покрывают прогоны небольшой команды",
        ],
        [
            "Машина под runner / agent",
            "ноутбук, облачная VM, self-hosted runner",
            "Облачная Linux VM 4 vCPU / 8 ГБ в Docker: прогон не зависит от ноутбука QA, "
            "стенд воспроизводим, браузер ставится в контейнер",
        ],
        [
            "Запуск тестов параллельно",
            "один поток, matrix-джобы, pytest-xdist",
            "Матрица CI по командам и тегам плюс параллельные воркеры Playwright: "
            "каждая команда — отдельная джоба, регресс не растягивается на часы",
        ],
        [
            "Система отчётов",
            "Allure server, GitHub Pages, артефакт CI",
            "HTML-отчёт как артефакт CI: ссылку получают QA и владельцы продуктов, "
            "история прогонов хранится в GitHub Actions",
        ],
    ],
    widths=[Cm(3.4), Cm(3.6), Cm(8.1)],
)

h("3.2. Схема: от репозитория до отчёта", 2)
img("schema.png", "Путь тестов: репозиторий → CI → runner → результаты → HTML-отчёт", width=15.5)
p("Простыми словами: тестировщик отправляет тесты в репозиторий (1), сервер CI замечает изменение и запускает "
  "проверку (2). На отдельной машине (3) тесты открывают браузер и работают параллельно (4). "
  "Каждый шаг записывается в файл результатов (5), из него собирается наглядный отчёт (6), "
  "а ссылка на отчёт уходит команде и на демо (7).")

h("3.3. Назначение файлов и папок проекта на behave", 2)
table(
    "Файлы и папки проекта",
    ["Путь", "Назначение"],
    [
        ["features/", "Сценарии на Gherkin (.feature): что проверяем — на языке бизнеса, без кода"],
        ["features/steps/", "Реализация шагов на Python: один шаг используется во многих сценариях"],
        ["features/environment.py", "Хуки before_all / before_scenario / after_*: запуск браузера, настройка data-test-id"],
        ["requirements.txt", "Список зависимостей: одной командой pip install -r разворачивается окружение"],
        ["reports/", "behave.json с результатами прогона и собранный из него behave-report.html; в git не коммитится"],
        [".github/workflows/tests.yml", "CI-пайплайн: установка зависимостей, прогон тестов, публикация отчёта"],
    ],
    widths=[Cm(5.0), Cm(10.0)],
)

# ---------- Задание 4 ----------
h("Задание 4. Якоря на UI и сценарий Gherkin", 1)

h("4.1. Атрибуты data-test-id на странице", 2)
p("В файл transfer.html добавлены атрибуты всем элементам, с которыми взаимодействует тест "
  "или в которых проверяется результат. Имена уникальны и говорят о назначении элемента.")
table(
    "Атрибуты data-test-id",
    ["Элемент", "data-test-id"],
    [
        ["Информация об остатке", "balance-info"],
        ["Список «счёт списания»", "transfer-source"],
        ["Список «счёт зачисления»", "transfer-target"],
        ["Поле суммы", "transfer-amount"],
        ["Кнопка «Перевести»", "transfer-submit"],
        ["Блок сообщения о результате", "transfer-message"],
    ],
    widths=[Cm(9.0), Cm(6.0)],
)

h("4.2. Сценарии в features/transfer.feature", 2)
p("Сценарии для успешного перевода и для ошибки при недостаточной сумме "
  "(плюс сценарий с нулевой суммой из задания 5.5).")
code((root / "features" / "transfer.feature").read_text(encoding="utf-8"), "features/transfer.feature")
img("form_success.png", "Успешный перевод 500 ₽", width=13.5)
img("form_error.png", "Ошибка при переводе 20000 ₽", width=13.5)

h("4.3. Почему тест на data-test-id устойчив", 2)
p("Локатор ищет элемент по атрибуту, а не по позиции в DOM, тексту или вложенности, поэтому перенос кнопки "
  "«Перевести» в другой блок или изменение вёрстки не меняет атрибут — тест продолжает находить тот же элемент.")
p("Правило обязательно для всех новых фич, потому что иначе локаторы начнут «плыть» при каждом рефакторинге: "
  "тесты будут падать без реальных дефектов, команда перестанет им доверять, а поддержка съест время, "
  "которое планировали на новое покрытие.")

# ---------- Задание 5 ----------
h("Задание 5. Шаги на Python, запуск и отчёт", 1)

h("5.1. features/environment.py", 2)
code((root / "features" / "environment.py").read_text(encoding="utf-8"),
     "запуск браузера и настройка поиска по data-test-id")

h("5.2. features/steps/transfer_steps.py", 2)
code((root / "features" / "steps" / "transfer_steps.py").read_text(encoding="utf-8"),
     "реализация шагов сценария")

h("5.3. Обращение в базу данных (db_check.py)", 2)
code((root / "db_check.py").read_text(encoding="utf-8"), "db_check.py")
p("Перевод выполняется в одной транзакции (BEGIN IMMEDIATE … COMMIT): если денег недостаточно, "
  "делается ROLLBACK и выбрасывается ValueError, поэтому остатки не меняются. Вывод программы: OK.")

h("5.4. Прогон тестов и отчёт", 2)
code(
    ".venv\\Scripts\\activate\n"
    "behave -f json -o reports/behave.json features\n"
    "python make_html_report.py reports/behave.json reports/behave-report.html",
    "команды запуска тестов и сборки отчёта",
)
img("behave_output.png", "Результат прогона: 1 фича, 3 сценария, 18 шагов — все прошли", width=15.5)
img("report_overview.png", "HTML-отчёт: сводка прогона и список сценариев", width=15.5)
img("report_steps.png", "HTML-отчёт: шаги сценариев", width=15.5)

h("5.5. Сценарий с нулевой суммой", 2)
p("Сценарий добавлен в transfer.feature и проходит без нового кода шагов, потому что шаги параметризованы: "
  "«вводит сумму {amount}» и «на экране отображается сообщение «{text}»» принимают любые данные, "
  "а правило «сумма должна быть больше нуля» уже реализовано в самой форме. Новый сценарий — это только "
  "новые данные, а не новый код.")

# ---------- Задание 6 ----------
h("Задание 6. Демо и план внедрения", 1)

h("6.1. Пять пунктов демо и порядок показа", 2)
table(
    "План демонстрации",
    ["№", "Что показываем", "Чью проблему это снимает"],
    [
        [
            "1",
            "Главная цель автоматизации: покрыть не менее половины регресса за 6 спринтов",
            "Владельцев продуктов — непонимание, зачем тратить бюджет на тесты",
        ],
        [
            "2",
            "Схема инфраструктуры: репозиторий → CI → runner → отчёт",
            "Разработчиков и аналитиков — вопрос «где это живёт и кто за это отвечает»",
        ],
        [
            "3",
            "Как создаются сценарии: Gherkin по-русски и data-test-id в разметке",
            "QA-инженеров — страх, что автоматизация «для программистов» и им придётся всё переписывать",
        ],
        [
            "4",
            "Видео 5–10 автотестов и HTML-отчёт после прогона",
            "Скептиков — недоверие «а это вообще работает и показывает правду»",
        ],
        [
            "5",
            "Эффект: сравнение ручного и автоматизированного прохождения, план внедрения по спринтам",
            "Руководства и закупки — вопрос «что дальше, сколько это стоит и когда окупится»",
        ],
    ],
    widths=[Cm(1.0), Cm(7.0), Cm(7.0)],
)

h("6.2. План на 6 спринтов для двух автоматизаторов", 2)
table(
    "План внедрения автоматизации",
    ["Спринт", "Автоматизаторы", "QA команд", "Что показываем на демо", "Покрытие"],
    [
        ["1", "Каркас: behave, Playwright, HTML-отчёт, CI; шаги формы перевода",
         "Реестр регресса, первые 10 кейсов «Личного кабинета»", "Каркас и 3 автотеста", "≈10 %"],
        ["2", "10 автотестов «Личного кабинета», проверки БД, параллельный запуск",
         "Сценарии Gherkin по своим кейсам, учатся запускать", "Отчёт и время прогона", "≈20 %"],
        ["3", "Приём тестов «Платежей», API-проверки через httpx",
         "Сценарии «Платежей», приоритизация денежных проверок", "API-тест перевода и остаток в БД", "≈30 %"],
        ["4", "Подключение «Карт», общие шаги, чистка нестабильных тестов",
         "Сценарии «Карт», ревью шагов", "0 нестабильных тестов за неделю", "≈40 %"],
        ["5", "Подключение «Кредитов», ночные прогоны, история отчётов",
         "Сценарии «Кредитов», дежурство по падениям", "Ночной прогон и динамика покрытия", "≈50 %"],
        ["6", "Подключение «Уведомлений», обучение QA, документация",
         "Сами пишут и поддерживают сценарии", "Итог: 55 % регресса, передача каркаса", "≈55 %"],
    ],
    widths=[Cm(1.2), Cm(5.2), Cm(4.3), Cm(3.6), Cm(1.7)],
)
p("Логика плана: одна новая команда в спринт, свободные дни автоматизаторы тратят на разработку шагов "
  "и на приём тестов, написанных QA. Первые два спринта — инвестиция в каркас, дальше покрытие растёт линейно.")

h("6.3. Первая команда и показатели", 2)
p("Первой подключается «Личный кабинет»: это frontend без сложных интеграций, ручных проверок много, "
  "а результат виден пользователю сразу — быстрый эффект и удобный пример для остальных команд. "
  "К тому же именно эта команда выбрана для демо.")
p("Показатели: (1) доля автоматизированного регресса, %; (2) время полного прогона и доля нестабильных тестов. "
  "Если покрытие не растёт — разбираем узкое место (нет шагов, нет данных, нет доступа к стенду) и снимаем его. "
  "Если растёт доля падений без дефектов — останавливаем набор новых тестов и стабилизируем существующие.")

# ---------- Контрольные вопросы ----------
doc.add_page_break()
h("Контрольные вопросы", 1)

qa = [
    (
        "1. Зачем перед выбором инструментов выяснять цель продукта и его потребителя?",
        "Инструмент выбирают под задачи и людей: если продукт денежный и критичный, нужны стабильные проверки "
        "API и БД, а не только UI; если тесты пишут ручные QA, важнее низкий порог входа, а не максимум "
        "возможностей. Без понимания цели можно взять дорогой стек, который команда не освоит.",
    ),
    (
        "2. Почему QA-инженеров команд называют клиентами инструмента автоматизации?",
        "Они пользуются результатом: читают отчёты, добавляют сценарии, принимают автотесты в свою работу. "
        "Если инструмент им неудобен, они вернутся к ручным проверкам, и вложения в автоматизацию не окупятся. "
        "Поэтому требования к каркасу собирают с них, как с заказчиков.",
    ),
    (
        "3. Какую задачу решает behave и чем он похож на Cucumber?",
        "Behave связывает сценарии на Gherkin (Дано / Когда / Тогда) с шагами на Python, поэтому тест читается "
        "как бизнес-описание. Как и Cucumber, он отделяет описание проверки от кода реализации и переиспользует "
        "шаги между сценариями.",
    ),
    (
        "4. Что такое data-test-id и почему он повышает надёжность автотестов?",
        "Это атрибут разметки, который существует только для тестов и не меняется при редизайне. "
        "Тест ищет элемент по нему, а не по позиции, тексту или CSS-классу, поэтому правки вёрстки "
        "не ломают локаторы и тесты падают только по реальным дефектам.",
    ),
    (
        "5. Зачем в environment.py вызывается set_test_id_attribute и что будет без него?",
        "Он говорит Playwright, что искать по test id нужно в атрибуте data-test-id, а не в принятом "
        "по умолчанию data-testid. Без него get_by_test_id ничего не найдёт и все сценарии упадут "
        "по таймауту локатора.",
    ),
    (
        "6. Чем отличаются before_all и before_scenario, почему страницу открывают заново?",
        "before_all выполняется один раз на весь прогон — там запускается браузер и задаётся адрес страницы. "
        "before_scenario выполняется перед каждым сценарием и открывает новую вкладку, чтобы сценарии "
        "не влияли друг на друга: состояние формы и сообщения не переходят из теста в тест.",
    ),
    (
        "7. Почему сценарий с нулевой суммой не потребовал нового кода шагов?",
        "Шаги параметризованы и уже покрывают любые сумму и текст сообщения, а проверка «сумма больше нуля» "
        "есть в самой форме. Новый сценарий — это новые данные для существующих шагов, поэтому код не менялся.",
    ),
    (
        "8. Зачем нужно демо и что на нём обязательно показывают?",
        "Демо нужно, чтобы заказчики и команды увидели работающий каркас и согласовали ожидания до того, "
        "как вложения станут большими. Обязательно показывают цель автоматизации, живой прогон тестов, "
        "отчёт о результатах и план внедрения с эффектом.",
    ),
    (
        "9. Почему на одного автоматизатора берут не больше пяти команд?",
        "Кроме написания тестов автоматизатор поддерживает каркас, разбирает падения, обучает QA и участвует "
        "в приёмке. При большем числе команд время уходит на переключение контекста, поддержка съедает весь "
        "ресурс, а новые тесты перестают появляться.",
    ),
    (
        "10. Какие показатели говорят, что внедрение работает, и что делать, если подход не работает?",
        "Работающий процесс виден по росту покрытия регресса, сокращению времени регресса и низкой доле "
        "падений без дефектов, плюс QA сами пишут сценарии. Если показатели не растут — ищем узкое место "
        "(данные, стенды, шаги, обучение), меняем подход и пересобираем план, а не наращиваем объём тестов.",
    ),
]
for question, answer in qa:
    p(question, bold=True)
    p(answer)

# ---------- Структура проекта ----------
h("Структура проекта", 1)
code(
    """transfer.html                     учебная форма перевода с data-test-id
features/transfer.feature          сценарии Gherkin
features/steps/transfer_steps.py   шаги на Python (Playwright)
features/environment.py            запуск браузера, настройка data-test-id
db_check.py                        проверка остатков через sqlite3
make_html_report.py                сборка HTML-отчёта из JSON-результатов behave
requirements.txt                   зависимости проекта
.github/workflows/tests.yml        CI-пайплайн
reports/                           behave.json и behave-report.html
screenshots/                       скриншоты для отчёта
tools/                             вспомогательные скрипты отчёта""",
    "файлы проекта",
)

h("Как запустить", 1)
code(
    """python -m venv .venv
.venv\\Scripts\\activate
pip install -r requirements.txt
playwright install chromium
behave -f json -o reports/behave.json features
python make_html_report.py reports/behave.json reports/behave-report.html
python db_check.py""",
    "команды полного цикла",
)
p("Фактический результат прогона: 1 фича, 3 сценария, 18 шагов — все пройдены, "
  "в HTML-отчёте 100 % успешных тестов; db_check.py выводит OK.", bold=True)

out = root / "Отчёт_Практическая_работа_Python_автоматизация_тестирования.docx"
doc.save(out)
print("saved", out.name, "| таблиц:", TABLES, "| рисунков:", FIGURES)
